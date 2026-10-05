"""Excel report: Summary, Clean data, Review.

Totals on the Summary sheet are live formulas over the Clean data sheet, so they always
agree with the data. Text that could run as a formula (starts with = + - @) is stored as text.
"""
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
SLATE, OCHRE_FILL, OCHRE_TEXT, MUTED = "2F4A63", "F6EFE0", "6B5320", "5C5850"
HEADER_FILL = PatternFill("solid", fgColor=SLATE)
WARN_FILL = PatternFill("solid", fgColor=OCHRE_FILL)
LINE = Border(bottom=Side(style="thin", color="DCD7CD"))
MONEY = "#,##0.00;(#,##0.00);-"
DANGEROUS_START = ("=", "+", "-", "@", "\t", "\r")
MAX_MONTHS = 120
ORDER_TEXT = {"month": "Month-first", "day": "Day-first"}
SOURCE_TEXT = {"chosen": "chosen by you", "detected": "detected from the file",
               "conflict": "the file mixes both orders, so the default was used",
               "assumed": "the file has no evidence either way, so the default was used"}


def font(**kw):
    return Font(name=FONT, size=kw.pop("size", 10), **kw)


def put(ws, row, col, value, **style):
    """Write a cell. Strings are always stored as text, never evaluated as formulas."""
    cell = ws.cell(row=row, column=col)
    if isinstance(value, str) and value.startswith(DANGEROUS_START):
        cell.value = value
        cell.data_type = "s"
        cell.quotePrefix = True
    else:
        cell.value = None if (value is not None and not isinstance(value, str) and pd.isna(value)) else value
    cell.font = style.get("font", font())
    if "fmt" in style:
        cell.number_format = style["fmt"]
    if "align" in style:
        cell.alignment = style["align"]
    if "fill" in style:
        cell.fill = style["fill"]
    return cell


def put_formula(ws, row, col, formula, **style):
    cell = ws.cell(row=row, column=col, value=formula)
    cell.font = style.get("font", font())
    cell.number_format = style.get("fmt", MONEY)
    return cell


def header_row(ws, row, labels):
    for i, label in enumerate(labels, start=1):
        put(ws, row, i, label, font=font(bold=True, color="FFFFFF"), fill=HEADER_FILL,
            align=Alignment(vertical="center", wrap_text=True))
    ws.row_dimensions[row].height = 22


# ------------------------------------------------------------------ Clean data
def write_clean_sheet(ws, cleaned: pd.DataFrame, mapping: dict) -> None:
    cols = list(cleaned.columns)
    header_row(ws, 1, cols)
    date_name = mapping["date"]
    for r, row in enumerate(cleaned.itertuples(index=False), start=2):
        for c, (name, value) in enumerate(zip(cols, row), start=1):
            if name == date_name:
                parsed = pd.to_datetime(value, errors="coerce") if isinstance(value, str) else pd.NaT
                put(ws, r, c, None if pd.isna(parsed) else parsed.date(), fmt="yyyy-mm-dd")
            elif name == mapping["amount"]:
                put(ws, r, c, None if pd.isna(value) else float(value), fmt=MONEY)
            elif name == "amount_outlier_flag":
                put(ws, r, c, "Yes" if value else "No")
            elif isinstance(value, (bool,)) or str(type(value)).endswith("bool'>"):
                put(ws, r, c, None if pd.isna(value) else ("Yes" if value else "No"))
            elif pd.api.types.is_scalar(value) and pd.isna(value):
                put(ws, r, c, None)
            else:
                put(ws, r, c, str(value) if not isinstance(value, (int, float)) else value)
    for i, name in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(32, len(name) + 6))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(cleaned) + 1}"


# ---------------------------------------------------------------------- Review
def review_rows(review: pd.DataFrame, info: pd.DataFrame, mapping: dict) -> list[list]:
    rows = []
    for rec in info.itertuples(index=False):
        rows.append([int(rec.source_row), rec.reason, rec.transaction_id, rec.date_original,
                     rec.amount_original, rec.merchant, rec.category])
    if len(review):
        lower = {c.lower(): c for c in review.columns}
        pick = lambda rec, name: rec.get(lower[name]) if name in lower else None  # noqa: E731
        for rec in review.to_dict("records"):
            note = ("first row kept" if rec["resolution"] == "kept"
                    else "removed from the clean data")
            rows.append([int(rec["data_row"]), f"Same transaction ID as another row; {note}",
                         pick(rec, "transaction_id"), rec.get(mapping["date"]),
                         rec.get(mapping["amount"]), pick(rec, "merchant"), pick(rec, "category")])
    return sorted(rows, key=lambda r: r[0])


def write_review_sheet(ws, rows: list[list]) -> None:
    header_row(ws, 1, ["Source row", "Why it's here", "Transaction ID", "Date (original)",
                       "Amount (original)", "Merchant", "Category"])
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            put(ws, r, c, value, align=Alignment(wrap_text=(c == 2), vertical="top"))
    for col, width in zip("ABCDEFG", (12, 52, 16, 18, 18, 24, 18)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:G{len(rows) + 1}"


# --------------------------------------------------------------------- Summary
def write_summary_sheet(ws, job: dict, cleaned: pd.DataFrame, mapping: dict) -> None:
    n = len(cleaned)
    cols = list(cleaned.columns)
    letter = lambda name: get_column_letter(cols.index(name) + 1)  # noqa: E731
    rng = lambda name: f"'Clean data'!${letter(name)}$2:${letter(name)}${n + 1}"  # noqa: E731
    amount, flag, when = rng(mapping["amount"]), rng("amount_outlier_flag"), rng(mapping["date"])
    cat_name = next((c for c in cols if c.lower() == "category"), None)

    for col, width in zip("ABC", (58, 18, 22)):
        ws.column_dimensions[col].width = width
    wrap = Alignment(wrap_text=True, vertical="top")

    put(ws, 1, 1, "Data quality report", font=font(size=16, bold=True))
    put(ws, 2, 1, job["input"]["filename"], font=font(color=MUTED))
    put(ws, 4, 1, "Quality score (out of 100)")
    put(ws, 4, 2, job["quality_score"], font=font(size=14, bold=True, color=SLATE), fmt="0.0")
    put(ws, 5, 1, "Rows in")
    put(ws, 5, 2, job["input"]["rows"], fmt="#,##0")
    put(ws, 6, 1, "Rows after cleaning")
    put_formula(ws, 6, 2, f"=ROWS({amount})", fmt="#,##0")
    put(ws, 7, 1, "Rules version")
    put(ws, 7, 2, job["rules_version"], align=Alignment(horizontal="right"))
    put(ws, 8, 1, "Slash dates read as")
    put(ws, 8, 2, ORDER_TEXT[job["date_order"]], align=Alignment(horizontal="right"))
    put(ws, 8, 3, SOURCE_TEXT[job["date_order_source"]], font=font(color=MUTED))

    row = 10
    ambiguous = next((i["count"] for i in job["issues"] if i["type"] == "ambiguous_dates"), 0)
    if ambiguous:
        other = "day-first" if job["date_order"] == "month" else "month-first"
        text = (f"{ambiguous:,} dates such as 03/04/23 could be read two ways. They were read "
                f"{ORDER_TEXT[job['date_order']].lower()} ({SOURCE_TEXT[job['date_order_source']]}). "
                f"If your file uses {other} dates, re-run with the date order switched.")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        put(ws, row, 1, text, font=font(color=OCHRE_TEXT), fill=WARN_FILL, align=wrap)
        ws.row_dimensions[row].height = 48
        row += 2

    put(ws, row, 1, "What we found", font=font(size=12, bold=True))
    row += 1
    header_row(ws, row, ["Issue", "Rows", ""])
    for issue in job["issues"]:
        row += 1
        put(ws, row, 1, issue["note"], align=wrap)
        put(ws, row, 2, issue["count"], fmt="#,##0")
    row += 1
    put(ws, row, 1, "Counts come from the cleaning pipeline (job.json). The full list of affected rows is on the Review sheet.",
        font=font(size=9, italic=True, color=MUTED))

    def totals_table(title, labels, label_fmt, all_f, flagged_f, other_label, total_all, total_flagged):
        nonlocal row
        row += 2
        put(ws, row, 1, title, font=font(size=12, bold=True))
        row += 1
        header_row(ws, row, [title.split(" by ")[-1].capitalize(), "All rows", "Excluding flagged"])
        first = row + 1
        for label in labels:
            row += 1
            put(ws, row, 1, label, fmt=label_fmt, align=Alignment(horizontal="left"))
            put_formula(ws, row, 2, all_f(row))
            put_formula(ws, row, 3, flagged_f(row))
        last = row
        row += 1
        put(ws, row, 1, other_label)
        put_formula(ws, row, 2, f"=B{row + 1}-SUM(B{first}:B{last})")
        put_formula(ws, row, 3, f"=C{row + 1}-SUM(C{first}:C{last})")
        row += 1
        put(ws, row, 1, "Total", font=font(bold=True))
        put_formula(ws, row, 2, total_all, font=font(bold=True))
        put_formula(ws, row, 3, total_flagged, font=font(bold=True))

    total_all, total_flagged = f"=SUM({amount})", f'=SUMIFS({amount},{flag},"No")'
    if cat_name:
        cats = sorted(cleaned[cat_name].dropna().astype(str).unique())
        cat = rng(cat_name)
        totals_table("Totals by category", cats, "General",
                     lambda r: f"=SUMIFS({amount},{cat},$A{r})",
                     lambda r: f'=SUMIFS({amount},{cat},$A{r},{flag},"No")',
                     "No category", total_all, total_flagged)

    dates = pd.to_datetime(cleaned[mapping["date"]], errors="coerce").dropna()
    if len(dates):
        months = pd.period_range(dates.min(), dates.max(), freq="M")
        if len(months) <= MAX_MONTHS:
            starts = [date(p.year, p.month, 1) for p in months]
            totals_table("Totals by month", starts, "mmm yyyy",
                         lambda r: f'=SUMIFS({amount},{when},">="&$A{r},{when},"<"&EDATE($A{r},1))',
                         lambda r: f'=SUMIFS({amount},{when},">="&$A{r},{when},"<"&EDATE($A{r},1),{flag},"No")',
                         "No date", total_all, total_flagged)
    row += 2
    put(ws, row, 1, "Totals are formulas over the Clean data sheet. Amounts are in the currency of the original file.",
        font=font(size=9, italic=True, color=MUTED))


def write_excel(path, job: dict, cleaned: pd.DataFrame, review: pd.DataFrame,
                info: pd.DataFrame, mapping: dict) -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    write_summary_sheet(summary, job, cleaned, mapping)
    write_clean_sheet(wb.create_sheet("Clean data"), cleaned, mapping)
    write_review_sheet(wb.create_sheet("Review"), review_rows(review, info, mapping))
    for sheet in wb.worksheets:  # print each sheet one page wide
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    wb.calculation.fullCalcOnLoad = True  # Excel recalculates every formula when the file opens
    wb.save(path)
