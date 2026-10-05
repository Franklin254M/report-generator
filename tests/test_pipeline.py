"""Run with:  python -m unittest discover -s tests   (also works under pytest)."""
import json
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from app.pipeline import main, run

ROOT = Path(__file__).resolve().parent.parent
SAMPLE = ROOT / "tests" / "messy_sample.csv"
GOLDEN = ROOT / "tests" / "golden_clean.csv"
REAL = ROOT / "samples" / "messy_finance_transactions.csv"


class Base(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def write_csv(self, name, text):
        path = self.tmp / name
        path.write_text(text, encoding="utf-8")
        return path


class PipelineTests(Base):
    def test_golden_output_matches_exactly(self):
        job = run(SAMPLE, self.tmp / "out")
        self.assertEqual(job["status"], "done")
        self.assertEqual((self.tmp / "out" / "clean.csv").read_text(), GOLDEN.read_text())

    def test_golden_job_facts(self):
        job = run(SAMPLE, self.tmp / "out")
        self.assertEqual(job["quality_score"], 60.9)
        counts = {i["type"]: i["count"] for i in job["issues"]}
        self.assertEqual(counts, {"duplicates": 1, "duplicate_ids": 1, "ambiguous_dates": 1,
                                  "missing_dates": 1, "unreadable_dates": 1, "missing_amounts": 1,
                                  "unreadable_amounts": 1, "outliers": 3})
        self.assertEqual(len(job["input"]["sha256"]), 64)
        self.assertEqual(job["rules_version"], "1.0")

    def test_same_id_different_content_goes_to_review_not_lost(self):
        out = self.tmp / "out"
        job = run(SAMPLE, out)
        review = (out / "review_duplicates.csv").read_text().splitlines()
        self.assertEqual(len(review), 3)  # header + kept + removed
        self.assertIn("kept", review[1])
        self.assertIn("removed", review[2])
        dedupe = job["stages"][1]
        self.assertEqual(dedupe["rows_in"], dedupe["rows_out"] + dedupe["exact_duplicates_removed"]
                         + dedupe["duplicate_ids_removed"])  # every row accounted for

    def test_ambiguous_dates_counted_and_order_switchable(self):
        month = run(SAMPLE, self.tmp / "m", date_order="month")
        run(SAMPLE, self.tmp / "d", date_order="day")
        self.assertEqual({i["type"]: i["count"] for i in month["issues"]}["ambiguous_dates"], 1)
        self.assertIn("T002,2023-01-02", (self.tmp / "m" / "clean.csv").read_text())
        self.assertIn("T002,2023-02-01", (self.tmp / "d" / "clean.csv").read_text())

    def test_missing_amount_column_asks_instead_of_crashing(self):
        path = self.write_csv("noamount.csv", "date,merchant,total\n2023-01-01,A,5\n")
        job = run(path, self.tmp / "out")
        self.assertEqual(job["status"], "needs_input")
        self.assertEqual(job["needs_input"]["field"], "amount_column")
        self.assertIn("total", job["needs_input"]["options"])
        self.assertEqual(main([str(path), "--out", str(self.tmp / "out2")]), 2)

    def test_rerun_with_column_mapping_succeeds(self):
        path = self.write_csv("noamount.csv", "date,merchant,total\n2023-01-01,A,5\n2023-01-02,B,7\n")
        job = run(path, self.tmp / "out", amount_column="total")
        self.assertEqual(job["status"], "done")
        self.assertEqual(job["quality_score"], 100.0)

    def test_empty_file_fails_with_readable_error(self):
        job = run(self.write_csv("empty.csv", ""), self.tmp / "out")
        self.assertEqual(job["status"], "failed")
        self.assertEqual(job["error"]["code"], "empty_file")
        self.assertTrue(job["error"]["fix"])

    def test_job_json_is_valid_and_written(self):
        run(SAMPLE, self.tmp / "out")
        data = json.loads((self.tmp / "out" / "job.json").read_text())
        self.assertEqual(data["status"], "done")
        self.assertFalse(list((self.tmp / "out").glob("*.tmp")))  # atomic write left no temp file

    @unittest.skipUnless(REAL.exists(), "real sample not present")
    def test_real_dataset_matches_the_original_tier1_numbers(self):
        job = run(REAL, self.tmp / "out")
        self.assertEqual(job["stages"][2]["rows_out"], 5000)
        counts = {i["type"]: i["count"] for i in job["issues"]}
        self.assertEqual(counts["duplicates"], 50)
        self.assertEqual(counts["outliers"], 227)


class ExcelAndDetectionTests(Base):
    """Step 2: Excel report, escaping, date-order detection."""

    def test_excel_has_three_sheets_and_live_formulas(self):
        run(SAMPLE, self.tmp / "out")
        wb = load_workbook(self.tmp / "out" / "report.xlsx")
        self.assertEqual(wb.sheetnames, ["Summary", "Clean data", "Review"])
        formulas = [c.value for row in wb["Summary"].iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.startswith("=")]
        self.assertTrue(any("SUMIFS" in f for f in formulas))
        self.assertEqual(wb["Clean data"].max_row, 22)  # header + 21 cleaned rows

    def test_text_that_looks_like_a_formula_is_stored_as_text(self):
        path = self.write_csv("evil.csv", "transaction_id,date,merchant,category,amount\n"
                              'T1,2023-01-01,"=HYPERLINK(""http://x"",""click"")",Shopping,5\n'
                              "T2,2023-01-02,+cmd,Shopping,6\nT3,2023-01-03,@SUM(1),Shopping,7\n"
                              "T4,2023-01-04,-2+3,Shopping,8\n")
        run(path, self.tmp / "out")
        ws = load_workbook(self.tmp / "out" / "report.xlsx")["Clean data"]
        merchants = [ws.cell(row=r, column=3) for r in range(2, 6)]
        self.assertTrue(all(c.data_type == "s" for c in merchants))
        self.assertTrue(merchants[0].value.startswith("=HYPERLINK"))
        self.assertFalse(any(c.data_type == "f" for row in ws.iter_rows() for c in row))

    def test_review_sheet_explains_every_problem_row(self):
        run(SAMPLE, self.tmp / "out")
        rows = list(load_workbook(self.tmp / "out" / "report.xlsx")["Review"].iter_rows(min_row=2, values_only=True))
        reasons = " | ".join(r[1] for r in rows)
        for expected in ("Date is missing", "Date couldn't be read", "Amount is missing",
                         "Amount couldn't be read", "Unusual amount", "Same transaction ID"):
            self.assertIn(expected, reasons)
        self.assertEqual(len(rows), 9)

    def test_date_order_is_detected_from_evidence(self):
        day = self.write_csv("day.csv", "date,amount\n25/12/2023,5\n03/04/23,6\n")
        job = run(day, self.tmp / "d")
        self.assertEqual((job["date_order"], job["date_order_source"]), ("day", "detected"))
        self.assertIn("2023-04-03", (self.tmp / "d" / "clean.csv").read_text())

        mixed = self.write_csv("mixed.csv", "date,amount\n25/12/2023,5\n12/25/2023,6\n")
        job = run(mixed, self.tmp / "m")
        self.assertEqual((job["date_order"], job["date_order_source"]), ("month", "conflict"))

        none = self.write_csv("none.csv", "date,amount\n03/04/23,5\n")
        job = run(none, self.tmp / "n")
        self.assertEqual((job["date_order"], job["date_order_source"]), ("month", "assumed"))

    def test_explicit_date_order_beats_detection(self):
        job = run(SAMPLE, self.tmp / "out", date_order="day")
        self.assertEqual(job["date_order_source"], "chosen")

    def test_excel_failure_still_delivers_csv(self):
        with mock.patch("app.pipeline.write_excel", side_effect=RuntimeError("boom")):
            job = run(SAMPLE, self.tmp / "out")
        self.assertEqual(job["status"], "done")
        self.assertTrue(job["warnings"])
        self.assertTrue((self.tmp / "out" / "clean.csv").exists())
        self.assertNotIn("report.xlsx", job["files"])


if __name__ == "__main__":
    unittest.main()
